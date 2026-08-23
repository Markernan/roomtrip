# -*- coding: utf-8 -*-
"""Genera docs/requerimientos-entrega.xlsx — versión que se entrega al profesor.

Contiene el catálogo de requerimientos y su clasificación, sin información de
seguimiento interno: no lleva estado de avance, ni cobertura de mockups, ni
responsables, ni el plan por laboratorio.

Uso:  python generar_entrega_xlsx.py [ruta_de_salida.xlsx]
"""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

import requerimientos_datos as D
from requerimientos_estilo import (
    F, BORDE, TURQ, TURQ_OSC, TURQ_MED, TURQ_PALE, BLANCO, BANDA_PAR, FUENTE,
    titulo_hoja, encabezado, banda, reglas_prioridad, pie_pagina)

OUT = sys.argv[1] if len(sys.argv) > 1 else \
    r"C:\Users\USUARIO\Desktop\roomtrip\docs\requerimientos-entrega.xlsx"

maestra = D.cargar()
wb = Workbook()

SUB = "1TEL05 Servicios y Aplicaciones para IoT · PUCP · Semestre 2026-2"

# ================================================================ Requerimientos
ws = wb.active
ws.title = "Requerimientos"
NC = 7
COLS = ["ID", "Módulo", "Rol", "Tipo", "Requerimiento", "Criterio de aceptación", "Prioridad"]
ANCHOS = [12, 24, 23, 13, 58, 66, 12]
titulo_hoja(ws, "RoomTrip · Lista de requerimientos", SUB, NC,
            "Sistema de gestión de reservas de alojamiento en hoteles vía aplicación móvil.")
encabezado(ws, 5, COLS, ANCHOS)

r = 6
for _, row in maestra.iterrows():
    for i, col in enumerate(COLS, start=1):
        c = ws.cell(row=r, column=i, value=row[col])
        c.font = F()
        c.border = BORDE
        centrada = i in (1, 3, 4, 7)
        c.alignment = Alignment(horizontal="center" if centrada else "left",
                                vertical="center" if centrada else "top", wrap_text=True)
        if r % 2 == 0:
            c.fill = PatternFill("solid", fgColor=BANDA_PAR)
    ws.cell(row=r, column=1).font = F(bold=True, color=TURQ_OSC)
    ws.row_dimensions[r].height = 40
    r += 1
ULT = r - 1

ws.auto_filter.ref = f"A5:G{ULT}"
ws.freeze_panes = "B6"
reglas_prioridad(ws, f"G6:G{ULT}")
pie_pagina(ws)

# ================================================================ Resumen
rs = wb.create_sheet("Resumen", 0)
NCR = 6
titulo_hoja(rs, "RoomTrip · Resumen del alcance", SUB, NCR,
            "Distribución de los requerimientos por rol, por módulo funcional y por tipo.")
for i, w in enumerate([30, 16, 16, 16, 16, 16], start=1):
    rs.column_dimensions[get_column_letter(i)].width = w

R = f"Requerimientos!$A$6:$A${ULT}"
ROL_R = f"Requerimientos!$C$6:$C${ULT}"
TIPO_R = f"Requerimientos!$D$6:$D${ULT}"
PRI_R = f"Requerimientos!$G$6:$G${ULT}"
MOD_R = f"Requerimientos!$B$6:$B${ULT}"


def kpi(fila, col, etiqueta, formula):
    c = rs.cell(row=fila, column=col, value=etiqueta)
    c.font = F(9, color=TURQ_OSC)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.fill = PatternFill("solid", fgColor=TURQ_PALE)
    c.border = BORDE
    v = rs.cell(row=fila + 1, column=col, value=formula)
    v.font = Font(name=FUENTE, size=20, bold=True, color=TURQ)
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.fill = PatternFill("solid", fgColor=TURQ_PALE)
    v.border = BORDE
    v.number_format = "0"


def rellena(fila, ncols):
    for col in range(1, ncols + 1):
        c = rs.cell(row=fila, column=col)
        c.border = BORDE
        if col > 1:
            c.font = F()
            c.alignment = Alignment(horizontal="center", vertical="center")
        if fila % 2 == 0:
            c.fill = PatternFill("solid", fgColor=BANDA_PAR)


banda(rs, 5, "Alcance del sistema", NCR)
rs.row_dimensions[6].height = 30
rs.row_dimensions[7].height = 36
kpi(6, 1, "Total de requerimientos", f"=COUNTA({R})")
kpi(6, 2, "Funcionales", f'=COUNTIF({TIPO_R},"Funcional")')
kpi(6, 3, "No funcionales", f'=COUNTIF({TIPO_R},"No funcional")')
kpi(6, 4, "Prioridad alta", f'=COUNTIF({PRI_R},"Alta")')
kpi(6, 5, "Prioridad media", f'=COUNTIF({PRI_R},"Media")')
kpi(6, 6, "Módulos funcionales", len(D.modulos(maestra)))

# --- por rol
banda(rs, 9, "Requerimientos por rol", NCR)
encabezado(rs, 10, ["Rol", "Total", "Funcionales", "No funcionales", "Alta", "Media"],
           [30, 16, 16, 16, 16, 16])
f0 = 11
for i, rol in enumerate(D.ORDEN_ROL):
    f = f0 + i
    rs.cell(row=f, column=1, value=rol).font = F(bold=True)
    rs.cell(row=f, column=2, value=f"=COUNTIF({ROL_R},$A{f})")
    rs.cell(row=f, column=3, value=f'=COUNTIFS({ROL_R},$A{f},{TIPO_R},"Funcional")')
    rs.cell(row=f, column=4, value=f'=COUNTIFS({ROL_R},$A{f},{TIPO_R},"No funcional")')
    rs.cell(row=f, column=5, value=f'=COUNTIFS({ROL_R},$A{f},{PRI_R},"Alta")')
    rs.cell(row=f, column=6, value=f'=COUNTIFS({ROL_R},$A{f},{PRI_R},"Media")')
    rellena(f, NCR)
f_tot = f0 + len(D.ORDEN_ROL)
rs.cell(row=f_tot, column=1, value="Total")
for col in range(2, NCR + 1):
    L = get_column_letter(col)
    rs.cell(row=f_tot, column=col, value=f"=SUM({L}{f0}:{L}{f_tot-1})")
for col in range(1, NCR + 1):
    c = rs.cell(row=f_tot, column=col)
    c.font = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=TURQ_MED)
    c.border = BORDE
    if col > 1:
        c.alignment = Alignment(horizontal="center", vertical="center")

# --- por modulo
f1 = f_tot + 2
banda(rs, f1, "Requerimientos por módulo funcional", NCR)
encabezado(rs, f1 + 1, ["Módulo", "Total", "Funcionales", "No funcionales", "Alta", "Media"],
           [30, 16, 16, 16, 16, 16])
mods = D.modulos(maestra)
f2 = f1 + 2
for i, mod in enumerate(mods):
    f = f2 + i
    rs.cell(row=f, column=1, value=mod).font = F()
    rs.cell(row=f, column=2, value=f"=COUNTIF({MOD_R},$A{f})")
    rs.cell(row=f, column=3, value=f'=COUNTIFS({MOD_R},$A{f},{TIPO_R},"Funcional")')
    rs.cell(row=f, column=4, value=f'=COUNTIFS({MOD_R},$A{f},{TIPO_R},"No funcional")')
    rs.cell(row=f, column=5, value=f'=COUNTIFS({MOD_R},$A{f},{PRI_R},"Alta")')
    rs.cell(row=f, column=6, value=f'=COUNTIFS({MOD_R},$A{f},{PRI_R},"Media")')
    rellena(f, NCR)
f3 = f2 + len(mods)
rs.cell(row=f3, column=1, value="Total").font = F(bold=True, color=BLANCO)
for col in range(2, NCR + 1):
    L = get_column_letter(col)
    rs.cell(row=f3, column=col, value=f"=SUM({L}{f2}:{L}{f3-1})")
for col in range(1, NCR + 1):
    c = rs.cell(row=f3, column=col)
    c.font = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=TURQ_MED)
    c.border = BORDE
    if col > 1:
        c.alignment = Alignment(horizontal="center", vertical="center")

ch = BarChart()
ch.type = "col"
ch.style = 2
ch.title = "Requerimientos por rol"
ch.y_axis.title = "Cantidad"
ch.height, ch.width = 7.5, 15
ch.add_data(Reference(rs, min_col=2, min_row=10, max_row=f_tot - 1), titles_from_data=True)
ch.set_categories(Reference(rs, min_col=1, min_row=f0, max_row=f_tot - 1))
ch.legend = None
rs.add_chart(ch, f"A{f3 + 2}")
rs.sheet_view.showGridLines = False
rs.freeze_panes = "A5"

# ================================================================ Por rol
pr = wb.create_sheet("Por rol", 2)
NCP = 6
titulo_hoja(pr, "RoomTrip · Requerimientos agrupados por rol", SUB, NCP,
            "Los mismos requerimientos de la hoja anterior, organizados por actor del sistema.")
encabezado(pr, 5, ["ID", "Requerimiento", "Criterio de aceptación", "Módulo", "Tipo", "Prioridad"],
           [12, 58, 64, 24, 13, 12])
f = 6
for rol in D.ORDEN_ROL:
    sub = maestra[maestra["Rol"] == rol]
    banda(pr, f, f"{rol}    ·    {len(sub)} {D.plural(len(sub))}", NCP)
    f += 1
    for _, row in sub.iterrows():
        for j, col in enumerate(["ID", "Requerimiento", "Criterio de aceptación",
                                 "Módulo", "Tipo", "Prioridad"], start=1):
            c = pr.cell(row=f, column=j, value=row[col])
            c.font = F(bold=True, color=TURQ_OSC) if j == 1 else F()
            c.border = BORDE
            centrada = j in (1, 4, 5, 6)
            c.alignment = Alignment(horizontal="center" if centrada else "left",
                                    vertical="center" if centrada else "top", wrap_text=True)
        pr.row_dimensions[f].height = 40
        f += 1
    f += 1
ULTP = f - 2
reglas_prioridad(pr, f"F6:F{ULTP}")
pr.freeze_panes = "A6"
pie_pagina(pr)

# ================================================================ Portada
po = wb.create_sheet("Portada")
NCG = 4
titulo_hoja(po, "RoomTrip", "Sistema de gestión de reservas de alojamiento en hoteles", NCG)
for i, w in enumerate([32, 60, 30, 26], start=1):
    po.column_dimensions[get_column_letter(i)].width = w

AMARILLO = PatternFill("solid", fgColor="FFF9C4")


def fila_dato(fila, etiqueta, valor, completar=False):
    c = po.cell(row=fila, column=1, value=etiqueta)
    c.font = F(bold=True, color=TURQ_OSC)
    c.border = BORDE
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    v = po.cell(row=fila, column=2, value=valor)
    v.font = F()
    v.border = BORDE
    v.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    if completar:
        v.fill = AMARILLO
    for col in (3, 4):
        po.cell(row=fila, column=col).border = BORDE
    po.row_dimensions[fila].height = 24
    return fila + 1


f = 5
banda(po, f, "Datos de la entrega", NCG)
f += 1
f = fila_dato(f, "Curso", "1TEL05 — Servicios y Aplicaciones para IoT")
f = fila_dato(f, "Carrera", "Ingeniería de las Telecomunicaciones — PUCP")
f = fila_dato(f, "Semestre", "2026-2")
f = fila_dato(f, "Proyecto", "RoomTrip")
f = fila_dato(f, "Entregable", "Lista de requerimientos")
f = fila_dato(f, "Fecha", dt.date.today().strftime("%d/%m/%Y"))
f = fila_dato(f, "Integrantes", "", completar=True)

po.merge_cells(start_row=f, start_column=1, end_row=f, end_column=NCG)
c = po.cell(row=f, column=1,
            value="Las celdas resaltadas en amarillo son las que debe completar el equipo "
                  "antes de entregar.")
c.font = F(9, italic=True, color="827717")
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
f += 2

banda(po, f, "Descripción del sistema", NCG)
f += 1
for texto in [
    "RoomTrip permite a un cliente consultar hoteles, habitaciones y servicios, reservar con "
    "tarjeta, comunicarse con el hotel por chat y solicitar un taxi gratuito al aeropuerto al "
    "hacer el checkout.",
    "El sistema se compone de dos aplicaciones con persistencias separadas: una app móvil "
    "Android nativa en Java, usada por los cuatro roles, y una app web de gestión de taxistas "
    "que expone una API REST.",
    "La app móvil no accede en ningún caso a la base de datos del sistema de taxistas: toda "
    "esa información se obtiene a través de la API REST.",
]:
    po.merge_cells(start_row=f, start_column=1, end_row=f, end_column=NCG)
    c = po.cell(row=f, column=1, value=texto)
    c.font = F()
    c.alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)
    po.row_dimensions[f].height = 34
    f += 1
f += 1

banda(po, f, "Contenido del libro", NCG)
f += 1
encabezado(po, f, ["Hoja", "Contenido", "", ""], [32, 60, 30, 26])
f += 1
for nombre, desc in [
    ("Resumen", "Distribución de los requerimientos por rol, por módulo funcional y por tipo."),
    ("Requerimientos", "Catálogo completo con su criterio de aceptación y prioridad."),
    ("Por rol", "Los mismos requerimientos agrupados por actor del sistema."),
]:
    po.cell(row=f, column=1, value=nombre).font = F(bold=True, color=TURQ_OSC)
    po.cell(row=f, column=2, value=desc).font = F()
    for col in range(1, NCG + 1):
        c = po.cell(row=f, column=col)
        c.border = BORDE
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    po.row_dimensions[f].height = 28
    f += 1
f += 1

banda(po, f, "Convenciones de identificación", NCG)
f += 1
for texto in [
    "RF-SA-nn · requerimiento funcional del Superadmin.",
    "RF-AH-nn · requerimiento funcional del Administrador de hotel.",
    "RF-TX-nn · requerimiento funcional del Taxista.",
    "RF-CL-nn · requerimiento funcional del Cliente.",
    "RF-GN-nn · requerimiento funcional transversal a todos los roles.",
    "RNF-nn · requerimiento no funcional del sistema.",
]:
    po.merge_cells(start_row=f, start_column=1, end_row=f, end_column=NCG)
    c = po.cell(row=f, column=1, value="•   " + texto)
    c.font = F()
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    po.row_dimensions[f].height = 20
    f += 1

po.sheet_view.showGridLines = False
wb.move_sheet("Portada", offset=-3)
wb.active = 0
wb.save(OUT)

print("Generado:", OUT)
print("Total requerimientos:", len(maestra))
print("Hojas:", wb.sheetnames)
print("Columnas de la hoja Requerimientos:", COLS)
