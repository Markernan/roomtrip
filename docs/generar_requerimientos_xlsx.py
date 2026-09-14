# -*- coding: utf-8 -*-
"""Genera docs/requerimientos.xlsx — versión interna del equipo.

Incluye estado de avance, cobertura de mockups y el plan por laboratorio.
NO es la versión que se entrega al profesor: para eso está generar_entrega_xlsx.py.

Uso:  python generar_requerimientos_xlsx.py [ruta_de_salida.xlsx]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule
from openpyxl.chart import BarChart, Reference

import requerimientos_datos as D
from requerimientos_estilo import (
    F, BORDE, TURQ, TURQ_OSC, TURQ_MED, TURQ_PALE, BLANCO, ROJO_T, GRIS_T, BANDA_PAR,
    FUENTE, titulo_hoja, encabezado, banda, reglas_estado, reglas_mockup,
    reglas_prioridad, pie_pagina)

OUT = sys.argv[1] if len(sys.argv) > 1 else \
    r"C:\Users\USUARIO\Desktop\roomtrip\docs\requerimientos.xlsx"

maestra = D.cargar()
wb = Workbook()

# ================================================================ Requerimientos
ws = wb.active
ws.title = "Requerimientos"
NC = 13
COLS = ["ID", "Módulo", "Rol", "Tipo", "Requerimiento", "Criterio de aceptación",
        "Prioridad", "Fases", "Cierre", "Estado", "Mockup", "Responsable", "Observaciones"]
ANCHOS = [11, 22, 21, 12, 52, 58, 10, 20, 12, 13, 10, 16, 46]
titulo_hoja(ws, "RoomTrip · Matriz de requerimientos",
            "1TEL05 Servicios y Aplicaciones para IoT · PUCP · Semestre 2026-2", NC,
            "Única hoja editable. «Fases» son los laboratorios en los que se trabaja el "
            "requerimiento; «Cierre» es aquel en el que queda terminado.")
encabezado(ws, 5, COLS, ANCHOS)

r = 6
for _, row in maestra.iterrows():
    for i, col in enumerate(COLS, start=1):
        c = ws.cell(row=r, column=i, value=row[col])
        c.font = F()
        c.border = BORDE
        c.alignment = Alignment(
            horizontal="center" if i in (1, 3, 4, 7, 8, 9, 10, 11) else "left",
            vertical="center" if i in (1, 3, 4, 7, 8, 9, 10, 11) else "top",
            wrap_text=True)
    ws.cell(row=r, column=1).font = F(bold=True, color=TURQ_OSC)
    ws.cell(row=r, column=8).font = F(color=TURQ_OSC)
    ws.row_dimensions[r].height = 42
    r += 1
ULT = r - 1

ws.auto_filter.ref = f"A5:M{ULT}"
ws.freeze_panes = "B6"
reglas_estado(ws, f"J6:J{ULT}")
reglas_mockup(ws, f"K6:K{ULT}")
reglas_prioridad(ws, f"G6:G{ULT}")
for formula, rng in (('"Pendiente,En progreso,Hecho,Bloqueado"', f"J6:J{ULT}"),
                     ('"Alta,Media,Baja"', f"G6:G{ULT}"),
                     ('"Sí,Parcial,No,NA"', f"K6:K{ULT}")):
    dv = DataValidation(type="list", formula1=formula, allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(rng)
pie_pagina(ws)

# ================================================================ Resumen
rs = wb.create_sheet("Resumen", 0)
NCR = 9
titulo_hoja(rs, "RoomTrip · Resumen de requerimientos",
            "Todos los valores se calculan con fórmulas sobre la hoja Requerimientos", NCR,
            "Se actualiza solo al cambiar el estado de un requerimiento.")
for i, w in enumerate([28, 13, 13, 13, 14, 13, 13, 14, 14], start=1):
    rs.column_dimensions[get_column_letter(i)].width = w

R = f"Requerimientos!$A$6:$A${ULT}"
ROL_R = f"Requerimientos!$C$6:$C${ULT}"
TIPO_R = f"Requerimientos!$D$6:$D${ULT}"
PRI_R = f"Requerimientos!$G$6:$G${ULT}"
FAS_R = f"Requerimientos!$H$6:$H${ULT}"
CIE_R = f"Requerimientos!$I$6:$I${ULT}"
EST_R = f"Requerimientos!$J$6:$J${ULT}"
MOC_R = f"Requerimientos!$K$6:$K${ULT}"
MOD_R = f"Requerimientos!$B$6:$B${ULT}"


def kpi(fila, col, etiqueta, formula, fmt="0", destaque=False):
    c = rs.cell(row=fila, column=col, value=etiqueta)
    c.font = F(9, color=TURQ_OSC)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.fill = PatternFill("solid", fgColor=TURQ_PALE)
    c.border = BORDE
    v = rs.cell(row=fila + 1, column=col, value=formula)
    v.font = Font(name=FUENTE, size=18, bold=True, color=(ROJO_T if destaque else TURQ))
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.fill = PatternFill("solid", fgColor=TURQ_PALE)
    v.border = BORDE
    v.number_format = fmt


def rellena(fila, ncols, pct_col=None):
    for col in range(1, ncols + 1):
        c = rs.cell(row=fila, column=col)
        c.border = BORDE
        if col > 1:
            c.font = F()
            c.alignment = Alignment(horizontal="center", vertical="center")
        if col == pct_col:
            c.number_format = "0.0%"
        if fila % 2 == 0:
            c.fill = PatternFill("solid", fgColor=BANDA_PAR)


banda(rs, 5, "Indicadores generales", NCR)
rs.row_dimensions[6].height = 30
rs.row_dimensions[7].height = 34
kpi(6, 1, "Total de requerimientos", f"=COUNTA({R})")
kpi(6, 2, "Funcionales", f'=COUNTIF({TIPO_R},"Funcional")')
kpi(6, 3, "No funcionales", f'=COUNTIF({TIPO_R},"No funcional")')
kpi(6, 4, "Prioridad alta", f'=COUNTIF({PRI_R},"Alta")')
kpi(6, 5, "Hechos", f'=COUNTIF({EST_R},"Hecho")')
kpi(6, 6, "En progreso", f'=COUNTIF({EST_R},"En progreso")')
kpi(6, 7, "Pendientes", f'=COUNTIF({EST_R},"Pendiente")')
kpi(6, 8, "% de avance", f'=IFERROR(COUNTIF({EST_R},"Hecho")/COUNTA({R}),0)', "0.0%")
kpi(6, 9, "Mockups faltantes", f'=COUNTIF({MOC_R},"No")', "0", destaque=True)

# --- por rol
banda(rs, 9, "Requerimientos por rol", NCR)
encabezado(rs, 10, ["Rol", "Total", "Alta", "Media", "Pendiente", "En progreso",
                    "Hecho", "% avance", "Sin mockup"],
           [28, 13, 13, 13, 14, 13, 13, 14, 14])
f0 = 11
for i, rol in enumerate(D.ORDEN_ROL):
    f = f0 + i
    rs.cell(row=f, column=1, value=rol).font = F(bold=True)
    rs.cell(row=f, column=2, value=f"=COUNTIF({ROL_R},$A{f})")
    rs.cell(row=f, column=3, value=f'=COUNTIFS({ROL_R},$A{f},{PRI_R},"Alta")')
    rs.cell(row=f, column=4, value=f'=COUNTIFS({ROL_R},$A{f},{PRI_R},"Media")')
    rs.cell(row=f, column=5, value=f'=COUNTIFS({ROL_R},$A{f},{EST_R},"Pendiente")')
    rs.cell(row=f, column=6, value=f'=COUNTIFS({ROL_R},$A{f},{EST_R},"En progreso")')
    rs.cell(row=f, column=7, value=f'=COUNTIFS({ROL_R},$A{f},{EST_R},"Hecho")')
    rs.cell(row=f, column=8, value=f"=IFERROR($G{f}/$B{f},0)")
    rs.cell(row=f, column=9, value=f'=COUNTIFS({ROL_R},$A{f},{MOC_R},"No")')
    rellena(f, NCR, pct_col=8)
f_tot = f0 + len(D.ORDEN_ROL)
rs.cell(row=f_tot, column=1, value="Total")
for col in range(2, 8):
    L = get_column_letter(col)
    rs.cell(row=f_tot, column=col, value=f"=SUM({L}{f0}:{L}{f_tot-1})")
rs.cell(row=f_tot, column=8, value=f"=IFERROR($G{f_tot}/$B{f_tot},0)")
rs.cell(row=f_tot, column=9, value=f"=SUM(I{f0}:I{f_tot-1})")
for col in range(1, NCR + 1):
    c = rs.cell(row=f_tot, column=col)
    c.font = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
    c.fill = PatternFill("solid", fgColor=TURQ_MED)
    c.border = BORDE
    if col > 1:
        c.alignment = Alignment(horizontal="center", vertical="center")
    if col == 8:
        c.number_format = "0.0%"
rs.conditional_formatting.add(f"H{f0}:H{f_tot-1}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color=TURQ_MED))

# --- por laboratorio
f1 = f_tot + 2
banda(rs, f1, "Carga por laboratorio", NCR)
rs.merge_cells(start_row=f1 + 1, start_column=1, end_row=f1 + 1, end_column=NCR)
c = rs.cell(row=f1 + 1, column=1,
            value="«Se trabaja» cuenta los requerimientos que se tocan en ese laboratorio; "
                  "«Cierra aquí», los que quedan terminados. Uno se trabaja en varios.")
c.font = F(9, italic=True, color=GRIS_T)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
encabezado(rs, f1 + 2, ["Laboratorio", "Fecha", "Días restantes", "Se trabaja",
                        "Cierra aquí", "Alta", "Pendiente", "% avance", "Sin mockup"],
           [28, 13, 13, 13, 14, 13, 13, 14, 14])
f2 = f1 + 3
for i, lab in enumerate(D.LABS):
    f = f2 + i
    rs.cell(row=f, column=1, value=D.LAB_NOMBRE[lab]).font = F(bold=True)
    fecha = D.LAB_FECHA[lab]
    if fecha:
        cf = rs.cell(row=f, column=2, value=fecha)
        cf.number_format = "DD/MM/YYYY"
        rs.cell(row=f, column=3, value=f"=$B{f}-TODAY()")
    else:
        rs.cell(row=f, column=2, value="Por definir")
        rs.cell(row=f, column=3, value="—")
    rs.cell(row=f, column=4, value=f'=COUNTIF({FAS_R},"*{lab}*")')
    rs.cell(row=f, column=5, value=f"=COUNTIF({CIE_R},$A{f})")
    rs.cell(row=f, column=6, value=f'=COUNTIFS({CIE_R},$A{f},{PRI_R},"Alta")')
    rs.cell(row=f, column=7, value=f'=COUNTIFS({CIE_R},$A{f},{EST_R},"Pendiente")')
    rs.cell(row=f, column=8, value=f'=IFERROR(COUNTIFS({CIE_R},$A{f},{EST_R},"Hecho")/$E{f},0)')
    rs.cell(row=f, column=9, value=f'=COUNTIFS({CIE_R},$A{f},{MOC_R},"No")')
    rellena(f, NCR, pct_col=8)
rs.conditional_formatting.add(f"H{f2}:H{f2+len(D.LABS)-1}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color=TURQ_MED))

# --- por modulo
f3 = f2 + len(D.LABS) + 1
banda(rs, f3, "Requerimientos por módulo funcional", NCR)
encabezado(rs, f3 + 1, ["Módulo", "Total", "Alta", "Pendiente", "Hecho", "% avance", "", "", ""],
           [28, 13, 13, 14, 13, 14, 13, 14, 14])
mods = D.modulos(maestra)
f4 = f3 + 2
for i, mod in enumerate(mods):
    f = f4 + i
    rs.cell(row=f, column=1, value=mod).font = F()
    rs.cell(row=f, column=2, value=f"=COUNTIF({MOD_R},$A{f})")
    rs.cell(row=f, column=3, value=f'=COUNTIFS({MOD_R},$A{f},{PRI_R},"Alta")')
    rs.cell(row=f, column=4, value=f'=COUNTIFS({MOD_R},$A{f},{EST_R},"Pendiente")')
    rs.cell(row=f, column=5, value=f'=COUNTIFS({MOD_R},$A{f},{EST_R},"Hecho")')
    rs.cell(row=f, column=6, value=f"=IFERROR($E{f}/$B{f},0)")
    rellena(f, 6, pct_col=6)
f5 = f4 + len(mods)

ch = BarChart()
ch.type = "col"
ch.style = 2
ch.title = "Requerimientos por rol"
ch.y_axis.title = "Cantidad"
ch.height, ch.width = 7.5, 16
ch.add_data(Reference(rs, min_col=2, min_row=10, max_row=f_tot - 1), titles_from_data=True)
ch.set_categories(Reference(rs, min_col=1, min_row=f0, max_row=f_tot - 1))
ch.legend = None
rs.add_chart(ch, f"A{f5 + 2}")
rs.sheet_view.showGridLines = False
rs.freeze_panes = "A5"


# ================================================================ vistas agrupadas
def vista(hoja, grupos, mapeo, ncols):
    f = 6
    rangos = []
    for etiqueta, sub in grupos:
        banda(hoja, f, etiqueta, ncols)
        ini = f + 1
        f += 1
        for _, row in sub.iterrows():
            c = hoja.cell(row=f, column=1, value=row["ID"])
            c.font = F(bold=True, color=TURQ_OSC)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDE
            for j, letra in enumerate(mapeo, start=2):
                cc = hoja.cell(row=f, column=j, value=(
                    f'=IFERROR(INDEX(Requerimientos!${letra}:${letra},'
                    f'MATCH($A{f},Requerimientos!$A:$A,0)),"")'))
                cc.font = F()
                cc.border = BORDE
                cc.alignment = Alignment(horizontal=("left" if j == 2 else "center"),
                                         vertical="center", wrap_text=True)
            hoja.row_dimensions[f].height = 30
            f += 1
        rangos.append((ini, f - 1))
        f += 1
    for ini, fin in rangos:
        hoja.row_dimensions.group(ini, fin, outline_level=1, hidden=False)
    return f - 2


pr = wb.create_sheet("Por rol", 2)
titulo_hoja(pr, "RoomTrip · Requerimientos agrupados por rol",
            "Vista de lectura. El contenido se trae por fórmula desde la hoja Requerimientos", 8,
            "Para modificar un requerimiento, edítalo en la hoja «Requerimientos».")
encabezado(pr, 5, ["ID", "Requerimiento", "Módulo", "Prioridad", "Cierre",
                   "Estado", "Mockup", "Responsable"], [11, 66, 22, 11, 12, 13, 10, 18])
grupos_rol = []
for rol in D.ORDEN_ROL:
    sub = maestra[maestra["Rol"] == rol]
    grupos_rol.append((f"{rol}    ·    {len(sub)} {D.plural(len(sub))}", sub))
ULTP = vista(pr, grupos_rol, ["E", "B", "G", "I", "J", "K", "L"], 8)
reglas_estado(pr, f"F6:F{ULTP}")
reglas_mockup(pr, f"G6:G{ULTP}")
reglas_prioridad(pr, f"D6:D{ULTP}")
pr.freeze_panes = "A6"
pie_pagina(pr)

pl = wb.create_sheet("Plan por laboratorio", 3)
titulo_hoja(pl, "RoomTrip · Plan por laboratorio",
            "Qué requerimientos se trabajan en cada entrega del curso", 8,
            "Un requerimiento aparece en varios laboratorios: son capas de madurez, no un "
            "reparto del alcance. La columna «Cierre» indica dónde queda terminado.")
encabezado(pl, 5, ["ID", "Requerimiento", "Rol", "Módulo", "Prioridad",
                   "Fases", "Cierre", "Estado"], [11, 62, 21, 22, 11, 20, 12, 13])
grupos_lab = []
for lab in D.LABS:
    sub = maestra[maestra["ID"].map(lambda i: lab in D.FASES[i])]
    fecha = D.LAB_FECHA[lab]
    txt = fecha.strftime("%d/%m/%Y") if fecha else "fecha por definir"
    cierran = int((maestra["Cierre"] == D.LAB_NOMBRE[lab]).sum())
    verbo = "cierra" if cierran == 1 else "cierran"
    grupos_lab.append((
        f"{D.LAB_NOMBRE[lab]}    ·    {txt}    ·    {D.LAB_DESC[lab]}    ·    "
        f"{len(sub)} {D.plural(len(sub))}, {cierran} {verbo} aquí", sub))
ULTL = vista(pl, grupos_lab, ["E", "C", "B", "G", "H", "I", "J"], 8)
reglas_estado(pl, f"H6:H{ULTL}")
reglas_prioridad(pl, f"E6:E{ULTL}")
pl.freeze_panes = "A6"
pie_pagina(pl)

# ================================================================ Guía
g = wb.create_sheet("Guía")
NCG = 4
titulo_hoja(g, "RoomTrip · Cómo usar este libro",
            "Matriz de requerimientos del proyecto · versión interna del equipo", NCG)
for i, w in enumerate([30, 62, 32, 24], start=1):
    g.column_dimensions[get_column_letter(i)].width = w


def tabla_texto(fila, titulo, cabeceras, filas, alto=30):
    banda(g, fila, titulo, NCG)
    fila += 1
    if cabeceras:
        encabezado(g, fila, cabeceras, [30, 62, 32, 24])
        fila += 1
    for datos in filas:
        for i, val in enumerate(datos, start=1):
            c = g.cell(row=fila, column=i, value=val)
            c.font = F(bold=True, color=TURQ_OSC) if i == 1 else F()
        for col in range(1, NCG + 1):
            cc = g.cell(row=fila, column=col)
            cc.border = BORDE
            cc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        g.row_dimensions[fila].height = alto
        fila += 1
    return fila + 1


def vinetas(fila, titulo, textos, alto=24):
    banda(g, fila, titulo, NCG)
    fila += 1
    for t in textos:
        g.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=NCG)
        c = g.cell(row=fila, column=1, value="•   " + t)
        c.font = F()
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        g.row_dimensions[fila].height = alto
        fila += 1
    return fila + 1


f = tabla_texto(5, "Hojas del libro", ["Hoja", "Para qué sirve", "¿Se edita?", ""], [
    ("Resumen", "Indicadores del proyecto: avance por rol, carga por laboratorio y por módulo.", "No. Se actualiza solo."),
    ("Requerimientos", "Tabla maestra con todos los requerimientos. Tiene filtros y listas desplegables.", "Sí. Es la única hoja editable."),
    ("Por rol", "Los mismos requerimientos agrupados por rol, para repartir el trabajo.", "No. Refleja la maestra."),
    ("Plan por laboratorio", "Qué se trabaja en cada entrega. Un requerimiento aparece en varias.", "No. Refleja la maestra."),
    ("Guía", "Esta hoja: significado de cada columna y convenciones.", "No."),
], alto=32)

f = vinetas(f, "Cómo se leen «Fases» y «Cierre»", [
    "Los laboratorios del curso no reparten el alcance del sistema: son capas de madurez sobre los mismos requerimientos.",
    "Ejemplo: «reservar con tarjeta» se dibuja en el Lab 3 como pantalla navegable y recién en el Lab 6 guarda la reserva de verdad. Sus fases son L3 · L6 y su cierre es el Lab 6.",
    "Por eso la suma de la columna «Se trabaja» del Resumen es mayor que el total. No es un error.",
    "Para saber qué se entrega en una fecha concreta, mira «Cierre», no «Fases».",
])

f = tabla_texto(f, "Significado de las columnas", ["Columna", "Qué contiene", "Valores admitidos", ""], [
    ("ID", "Identificador trazable con el enunciado del curso.", "RF-<ROL>-<NN> · RNF-<NN>"),
    ("Módulo", "Agrupación funcional del sistema.", "M1 a M10"),
    ("Rol", "Quién ejecuta el requerimiento.", "Los 4 roles + Transversal"),
    ("Tipo", "Naturaleza del requerimiento.", "Funcional · No funcional"),
    ("Requerimiento", "Enunciado en una línea.", "Texto libre"),
    ("Criterio de aceptación", "Cómo se comprueba que está cumplido.", "Texto libre"),
    ("Prioridad", "Urgencia relativa.", "Alta · Media · Baja"),
    ("Fases", "Laboratorios en los que se trabaja el requerimiento.", "L3 · L4 · L5 · L6 · L7 · Web"),
    ("Cierre", "Laboratorio en el que queda terminado.", "Lab 3 a Lab 7 · App web"),
    ("Estado", "Avance real.", "Pendiente · En progreso · Hecho · Bloqueado"),
    ("Mockup", "Si la pantalla ya está diseñada en Figma.", "Sí · Parcial · No · NA"),
    ("Responsable", "Integrante del equipo a cargo.", "Nombre"),
    ("Observaciones", "Notas, riesgos y detalles pendientes.", "Texto libre"),
], alto=28)

vinetas(f, "Convenciones de trabajo", [
    "Editar únicamente la hoja «Requerimientos». Las demás se actualizan solas.",
    "Al añadir un requerimiento, insertar la fila dentro del rango existente para que las fórmulas del Resumen lo tomen.",
    "Los IDs no se reutilizan: si un requerimiento se elimina, su ID queda retirado.",
    "Este libro es de uso interno. La versión que se entrega al profesor se genera con generar_entrega_xlsx.py.",
], alto=22)

g.sheet_view.showGridLines = False
wb.move_sheet("Guía", offset=-4)
wb.active = 1
wb.save(OUT)

print("Generado:", OUT)
print("Total requerimientos:", len(maestra))
print("Carga por laboratorio:")
for lab in D.LABS:
    trabaja = sum(1 for i in maestra["ID"] if lab in D.FASES[i])
    cierra = int((maestra["Cierre"] == D.LAB_NOMBRE[lab]).sum())
    print(f"  {D.LAB_NOMBRE[lab]:<10} se trabaja: {trabaja:<3} cierra: {cierra}")
