# -*- coding: utf-8 -*-
"""Estilos y helpers de dibujo compartidos por los generadores de Excel de RoomTrip."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

TURQ_OSC = "00696E"
TURQ = "00838F"
TURQ_MED = "4DB6AC"
TURQ_CLA = "B2DFDB"
TURQ_PALE = "E0F2F1"
GRIS_TXT = "37474F"
BLANCO = "FFFFFF"
VERDE_F, VERDE_T = "C8E6C9", "1B5E20"
AMBAR_F, AMBAR_T = "FFE0B2", "E65100"
ROJO_F, ROJO_T = "FFCDD2", "B71C1C"
GRIS_F, GRIS_T = "ECEFF1", "546E7A"
BANDA_PAR = "F7FBFB"

FUENTE = "Arial"
_thin = Side(style="thin", color="CFD8DC")
BORDE = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def F(size=10, bold=False, color=GRIS_TXT, italic=False):
    return Font(name=FUENTE, size=size, bold=bold, color=color, italic=italic)


def titulo_hoja(ws, texto, subtitulo, ncols, nota=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(name=FUENTE, size=16, bold=True, color=BLANCO)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=1, value=subtitulo)
    c.font = Font(name=FUENTE, size=9, color=BLANCO)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=TURQ)
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=TURQ_MED)
    if nota:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        c = ws.cell(row=3, column=1, value=nota)
        c.font = Font(name=FUENTE, size=9, italic=True, color=TURQ_OSC)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[3].height = 16


def encabezado(ws, fila, headers, anchos):
    for i, (h, w) in enumerate(zip(headers, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=TURQ_OSC)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 30


def banda(ws, fila, texto, ncols, fill=TURQ_CLA, color=TURQ_OSC, size=11):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    for col in range(1, ncols + 1):
        cc = ws.cell(row=fila, column=col)
        cc.fill = PatternFill("solid", fgColor=fill)
        cc.border = BORDE
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(name=FUENTE, size=size, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[fila].height = 22


def reglas_estado(ws, rango):
    for val, fill, txt, bold in (("Hecho", VERDE_F, VERDE_T, True),
                                 ("En progreso", AMBAR_F, AMBAR_T, True),
                                 ("Bloqueado", ROJO_F, ROJO_T, True),
                                 ("Pendiente", GRIS_F, GRIS_T, False)):
        ws.conditional_formatting.add(rango, CellIsRule(
            operator="equal", formula=[f'"{val}"'],
            fill=PatternFill("solid", fgColor=fill),
            font=Font(name=FUENTE, size=10, bold=bold, color=txt)))


def reglas_mockup(ws, rango):
    for val, fill, txt, bold in (("Sí", VERDE_F, VERDE_T, False),
                                 ("Parcial", AMBAR_F, AMBAR_T, False),
                                 ("No", ROJO_F, ROJO_T, True)):
        ws.conditional_formatting.add(rango, CellIsRule(
            operator="equal", formula=[f'"{val}"'],
            fill=PatternFill("solid", fgColor=fill),
            font=Font(name=FUENTE, size=10, bold=bold, color=txt)))


def reglas_prioridad(ws, rango):
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="equal", formula=['"Alta"'],
        font=Font(name=FUENTE, size=10, bold=True, color=ROJO_T)))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="equal", formula=['"Media"'],
        font=Font(name=FUENTE, size=10, color=AMBAR_T)))


def pie_pagina(ws, fila_titulo=5, horizontal=True):
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = f"{fila_titulo}:{fila_titulo}"
    if horizontal:
        ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
